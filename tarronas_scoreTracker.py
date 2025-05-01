import os
import openpyxl
from openpyxl import Workbook, load_workbook
import tkinter as tk
from tkinter import messagebox, ttk

FILENAME = "student_scores.xlsx"

# Create workbook if it doesn't exist
def init_workbook():
    if not os.path.exists(FILENAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Scores"
        ws.append(["Name", "Score", "Status"])
        wb.save(FILENAME)

# Determine pass/fail
def get_status(score):
    return "Pass" if score >= 50 else "Fail"

# Add or update student record
def add_or_update_student(name, score):
    try:
        score = int(score)
    except ValueError:
        messagebox.showerror("Invalid Input", "Score must be a number.")
        return

    status = get_status(score)
    wb = load_workbook(FILENAME)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            row[1].value = score
            row[2].value = status
            wb.save(FILENAME)
            messagebox.showinfo("Updated", f"{name}'s score updated.")
            return

    ws.append([name, score, status])
    wb.save(FILENAME)
    messagebox.showinfo("Added", f"{name} added successfully.")

# Display all records in the GUI
def display_all_records():
    wb = load_workbook(FILENAME)
    ws = wb.active
    records_list.delete(*records_list.get_children())
    for row in ws.iter_rows(min_row=2, values_only=True):
        records_list.insert("", "end", values=row)

# Submit button callback
def submit_record():
    name = name_entry.get().strip()
    score = score_entry.get().strip()
    if not name or not score:
        messagebox.showwarning("Missing Data", "Both fields are required.")
        return
    add_or_update_student(name, score)
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)
    display_all_records()

# Initialize Excel file
init_workbook()

# Tkinter GUI setup
root = tk.Tk()
root.title("Student Score Tracker")
root.geometry("500x400")

# Input frame
input_frame = tk.Frame(root)
input_frame.pack(pady=10)

tk.Label(input_frame, text="Student Name:").grid(row=0, column=0, padx=5, pady=5, sticky="e")
name_entry = tk.Entry(input_frame, width=25)
name_entry.grid(row=0, column=1, padx=5, pady=5)

tk.Label(input_frame, text="Score:").grid(row=1, column=0, padx=5, pady=5, sticky="e")
score_entry = tk.Entry(input_frame, width=25)
score_entry.grid(row=1, column=1, padx=5, pady=5)

submit_btn = tk.Button(root, text="Add/Update Record", command=submit_record)
submit_btn.pack(pady=10)

# Records display
records_frame = tk.Frame(root)
records_frame.pack(fill="both", expand=True)

columns = ("Name", "Score", "Status")
records_list = ttk.Treeview(records_frame, columns=columns, show="headings")
for col in columns:
    records_list.heading(col, text=col)
    records_list.column(col, anchor="center")

records_list.pack(fill="both", expand=True, padx=10, pady=10)

# Load existing records initially
display_all_records()

root.mainloop()
